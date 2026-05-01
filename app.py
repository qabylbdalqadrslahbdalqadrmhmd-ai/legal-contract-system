import subprocess
import sys

def install(pkg):
    subprocess.check_call([sys.executable, "-m", "pip", "install", pkg, "-q"])

try:
    import google.generativeai as genai
except ImportError
    install("google-generativeai")
    import google.generativeai as genai

import subprocess
import sys

def install(pkg):
    subprocess.check_call([sys.executable, "-m", "pip", "install", pkg, "-q"])

try:
    import google.generativeai as genai
except ImportError:
    install("google-generativeai")
    import google.generativeai as genai

try:
    import fitz
except ImportError:
    install("pymupdf")
    import fitz

try:
    from docx import Document
except ImportError:
    install("python-docx")
    from docx import Document

import streamlit as st
import base64
import json
import re
import io

PYMUPDF_AVAILABLE = True
DOCX_AVAILABLE = True

# ─────────────────────────────────────────────
# Page config
# ─────────────────────────────────────────────
st.set_page_config(
    page_title="محلل العقود الذكي",
    page_icon="⚖️",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ─────────────────────────────────────────────
# Custom CSS
# ─────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Tajawal:wght@300;400;500;700;800&display=swap');

* { font-family: 'Tajawal', sans-serif !important; }

html, body, [class*="css"] {
    direction: rtl;
    background: #0f0f1a;
    color: #e8e8f0;
}

.stApp { background: #0f0f1a; }

.main-header {
    background: linear-gradient(135deg, #1a1a2e 0%, #16213e 50%, #0f3460 100%);
    border: 1px solid #e94560;
    border-radius: 16px;
    padding: 2rem;
    text-align: center;
    margin-bottom: 2rem;
}
.main-header h1 { color: #e94560; font-size: 2.5rem; font-weight: 800; margin: 0; text-shadow: 0 0 30px rgba(233,69,96,0.5); }
.main-header p  { color: #a0a0c0; font-size: 1.1rem; margin: 0.5rem 0 0; }

.analysis-card {
    background: linear-gradient(135deg, #1a1a2e, #16213e);
    border: 1px solid #2a2a4a;
    border-radius: 12px;
    padding: 1.5rem;
    margin-bottom: 1.2rem;
}
.strength-card  { border-right: 4px solid #00d4aa; background: linear-gradient(135deg,#0d2b25,#0f2d27); }
.weakness-card  { border-right: 4px solid #e94560; background: linear-gradient(135deg,#2b0d1a,#2d0f1e); }
.suggestion-card{ border-right: 4px solid #f7c948; background: linear-gradient(135deg,#2b2500,#2d2600); }
.missing-card   { border-right: 4px solid #ff6b35; background: linear-gradient(135deg,#2b1500,#2d1600); }

.card-title { font-size: 1.2rem; font-weight: 700; margin-bottom: 1rem; display: flex; align-items: center; gap: 0.5rem; }
.strength-title  { color: #00d4aa; }
.weakness-title  { color: #e94560; }
.suggestion-title{ color: #f7c948; }
.missing-title   { color: #ff6b35; }

.score-ring {
    width: 120px; height: 120px; border-radius: 50%;
    display: flex; align-items: center; justify-content: center;
    margin: 0 auto 1rem; font-size: 2rem; font-weight: 800; border: 6px solid;
}
.score-high { border-color:#00d4aa; color:#00d4aa; background:rgba(0,212,170,0.1); }
.score-mid  { border-color:#f7c948; color:#f7c948; background:rgba(247,201,72,0.1); }
.score-low  { border-color:#e94560; color:#e94560; background:rgba(233,69,96,0.1); }

.stTabs [data-baseweb="tab-list"] { background:#1a1a2e; border-radius:10px; padding:4px; gap:4px; }
.stTabs [data-baseweb="tab"] { background:transparent; border-radius:8px; color:#a0a0c0; font-weight:600; padding:8px 20px; }
.stTabs [aria-selected="true"] { background:#e94560 !important; color:white !important; }

.stButton > button {
    background: linear-gradient(135deg,#e94560,#c73652);
    color: white; border: none; border-radius: 10px;
    padding: 0.7rem 2rem; font-size: 1.1rem; font-weight: 700; width: 100%;
}
.stButton > button:hover { transform: translateY(-2px); box-shadow: 0 8px 25px rgba(233,69,96,0.4); }

.stTextArea textarea, .stTextInput input {
    background:#1a1a2e !important; border:1px solid #2a2a4a !important;
    border-radius:10px !important; color:#e8e8f0 !important; direction:rtl;
}
.stTextArea textarea:focus, .stTextInput input:focus {
    border-color:#e94560 !important; box-shadow:0 0 0 2px rgba(233,69,96,0.2) !important;
}

[data-testid="stSidebar"] { background:#12121f !important; border-left:1px solid #2a2a4a; }

.tag { display:inline-block; padding:3px 12px; border-radius:20px; font-size:0.85rem; font-weight:600; margin:3px; }
.tag-green  { background:rgba(0,212,170,0.2); color:#00d4aa; border:1px solid #00d4aa; }
.tag-red    { background:rgba(233,69,96,0.2);  color:#e94560; border:1px solid #e94560; }
.tag-yellow { background:rgba(247,201,72,0.2); color:#f7c948; border:1px solid #f7c948; }

.section-divider { height:2px; background:linear-gradient(90deg,#e94560,transparent); margin:1.5rem 0; border-radius:2px; }

.contract-output {
    background:#1a1a2e; border:1px solid #2a2a4a; border-radius:12px;
    padding:2rem; white-space:pre-wrap; line-height:2; direction:rtl;
}
</style>
""", unsafe_allow_html=True)

# ─────────────────────────────────────────────
# Header
# ─────────────────────────────────────────────
st.markdown("""
<div class="main-header">
    <h1>⚖️ محلل العقود الذكي</h1>
    <p>تحليل شامل واحترافي للعقود القانونية بالذكاء الاصطناعي — مجاني 100%</p>
</div>
""", unsafe_allow_html=True)

# ─────────────────────────────────────────────
# Sidebar
# ─────────────────────────────────────────────
with st.sidebar:
    st.markdown("### 🔑 إعدادات API")
    st.markdown("""
<div style="background:rgba(0,212,170,0.1);border:1px solid #00d4aa;border-radius:8px;padding:0.8rem;margin-bottom:1rem;font-size:0.85rem;color:#00d4aa">
    ✅ Google Gemini مجاني تماماً!<br>
    احصل على مفتاحك من:<br>
    <a href="https://aistudio.google.com" target="_blank" style="color:#00d4aa">aistudio.google.com</a>
</div>
""", unsafe_allow_html=True)
    api_key = st.text_input("Gemini API Key", type="password", placeholder="AIza...")

    st.markdown("---")
    st.markdown("### 📋 نوع العقد")
    contract_type = st.selectbox("اختر نوع العقد", [
        "تلقائي (الذكاء يحدده)", "عقد عمل", "عقد إيجار", "عقد بيع",
        "عقد شراكة", "عقد خدمات", "عقد توريد", "عقد استشارات",
        "عقد نقل ملكية", "عقد قرض", "أخرى",
    ])

    st.markdown("---")
    st.markdown("### ⚙️ خيارات التحليل")
    check_missing        = st.checkbox("🔍 تحقق من البنود الناقصة", value=True)
    check_risks          = st.checkbox("⚠️ تحليل المخاطر القانونية", value=True)
    suggest_improvements = st.checkbox("💡 اقتراح تحسينات", value=True)
    generate_score       = st.checkbox("🎯 تقييم شامل (نقاط)", value=True)

    st.markdown("---")
    st.markdown("### 📊 مستوى التفصيل")
    detail_level = st.radio("", ["مختصر", "متوسط", "تفصيلي مكثف"], index=1)


# ─────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────
def extract_text_from_pdf(file_bytes: bytes) -> str:
    if not PYMUPDF_AVAILABLE:
        return ""
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    return "\n".join(page.get_text("text") for page in doc).strip()

def extract_text_from_docx(file_bytes: bytes) -> str:
    if not DOCX_AVAILABLE:
        return ""
    doc = Document(io.BytesIO(file_bytes))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

def file_to_base64(file_bytes: bytes) -> str:
    return base64.standard_b64encode(file_bytes).decode("utf-8")

def score_class(s):
    return "score-high" if s >= 70 else ("score-mid" if s >= 50 else "score-low")


# ─────────────────────────────────────────────
# Gemini: Analyze contract
# ─────────────────────────────────────────────
def analyze_contract_gemini(api_key: str, contract_text: str,
                             contract_type_hint: str, options: dict,
                             image_data: dict | None = None) -> dict:

    genai.configure(api_key=api_key)
    model = genai.GenerativeModel("gemini-1.5-flash")

    detail_map = {"مختصر": "موجز", "متوسط": "متوسط التفصيل", "تفصيلي مكثف": "تفصيلي جداً"}
    detail = detail_map.get(options["detail_level"], "متوسط التفصيل")

    options_text = []
    if options["check_missing"]:        options_text.append("- تحقق من البنود الناقصة")
    if options["check_risks"]:          options_text.append("- حلل المخاطر القانونية")
    if options["suggest_improvements"]: options_text.append("- قدم اقتراحات تحسين")
    if options["generate_score"]:       options_text.append("- قيّم العقد بنقاط من 100")

    prompt = f"""أنت محامٍ وخبير قانوني متخصص في تحليل العقود، تتحدث العربية بطلاقة.
أجب فقط بـ JSON صحيح بدون أي نص إضافي أو ```json أو أي شيء آخر.

حلل العقد التالي وأعطني JSON بهذا التنسيق بالضبط:

{{
  "contract_type": "نوع العقد",
  "summary": "ملخص العقد في 3-5 جمل",
  "parties": ["الطرف الأول: ...", "الطرف الثاني: ..."],
  "key_terms": ["بند مهم 1", "بند مهم 2"],
  "strengths": [
    {{"title": "عنوان نقطة القوة", "detail": "شرح تفصيلي"}}
  ],
  "weaknesses": [
    {{"title": "عنوان نقطة الضعف", "detail": "شرح تفصيلي"}}
  ],
  "missing_clauses": [
    {{"title": "البند الناقص", "importance": "عالي", "suggestion": "نص مقترح"}}
  ],
  "risks": [
    {{"title": "الخطر", "severity": "عالي", "mitigation": "طريقة التخفيف"}}
  ],
  "improvements": [
    {{"title": "التحسين", "detail": "شرح التحسين", "priority": "عالي"}}
  ],
  "score": {{
    "overall": 75,
    "clarity": 80,
    "completeness": 70,
    "legal_protection": 75,
    "balance": 65,
    "justification": "سبب التقييم"
  }},
  "verdict": "الحكم النهائي على العقد",
  "recommendation": "توصية مفصلة للشخص قبل التوقيع"
}}

نوع العقد: {contract_type_hint}
مستوى التفصيل: {detail}
خيارات التحليل:
{chr(10).join(options_text)}

نص العقد:
{contract_text[:12000]}
"""

    if image_data:
        image_part = {
            "mime_type": image_data["media_type"],
            "data": base64.b64decode(image_data["data"])
        }
        response = model.generate_content([prompt, image_part])
    else:
        response = model.generate_content(prompt)

    raw = response.text.strip()
    raw = re.sub(r"^```(?:json)?\s*", "", raw)
    raw = re.sub(r"\s*```$", "", raw)
    return json.loads(raw)


# ─────────────────────────────────────────────
# Gemini: Generate contract
# ─────────────────────────────────────────────
def generate_contract_gemini(api_key: str, contract_type: str, details: str) -> str:
    genai.configure(api_key=api_key)
    model = genai.GenerativeModel("gemini-1.5-flash")

    prompt = f"""أنت محامٍ خبير في صياغة العقود القانونية باللغة العربية.
اكتب عقد {contract_type} احترافي وشامل باللغة العربية.

التفاصيل:
{details}

اكتب العقد بالكامل مع:
- ديباجة العقد وأطرافه
- جميع البنود الأساسية والفرعية
- الشروط والأحكام
- بنود الضمانات والتعويضات
- بنود فسخ العقد
- القانون الحاكم وحل النزاعات
- التواقيع والتوثيق

اجعل العقد متوازناً ويحمي حقوق الطرفين."""

    response = model.generate_content(prompt)
    return response.text


# ─────────────────────────────────────────────
# Main Tabs
# ─────────────────────────────────────────────
tab1, tab2, tab3 = st.tabs(["📄 تحليل عقد موجود", "✍️ إنشاء عقد جديد", "📚 دليل البنود"])

# ══════════════════════════════════════════════
# TAB 1 – ANALYZE
# ══════════════════════════════════════════════
with tab1:
    st.markdown("### 📤 رفع أو كتابة العقد")
    input_method = st.radio("طريقة الإدخال", ["📝 نص مباشر", "📎 PDF", "📄 Word", "🖼️ صورة"], horizontal=True)

    contract_text = ""
    image_data_dict = None

    if input_method == "📝 نص مباشر":
        contract_text = st.text_area("الصق نص العقد هنا", height=300,
                                     placeholder="اكتب أو الصق نص العقد هنا...")

    elif input_method == "📎 PDF":
        pdf_file = st.file_uploader("ارفع ملف PDF", type=["pdf"])
        if pdf_file:
            file_bytes = pdf_file.read()
            if PYMUPDF_AVAILABLE:
                with st.spinner("⏳ جاري استخراج النص..."):
                    contract_text = extract_text_from_pdf(file_bytes)
                if contract_text:
                    st.success(f"✅ تم استخراج {len(contract_text)} حرف")
                    with st.expander("👁️ عرض النص المستخرج"):
                        st.text(contract_text[:3000] + ("..." if len(contract_text) > 3000 else ""))
                else:
                    st.warning("⚠️ لم يُستخرج نص — سيُرسل كصورة")
                    image_data_dict = {"media_type": "application/pdf", "data": file_to_base64(file_bytes)}
            else:
                image_data_dict = {"media_type": "application/pdf", "data": file_to_base64(file_bytes)}
                st.info("📤 سيُرسل PDF للذكاء الاصطناعي مباشرة")

    elif input_method == "📄 Word":
        docx_file = st.file_uploader("ارفع ملف Word", type=["docx"])
        if docx_file:
            file_bytes = docx_file.read()
            if DOCX_AVAILABLE:
                with st.spinner("⏳ جاري استخراج النص..."):
                    contract_text = extract_text_from_docx(file_bytes)
                if contract_text:
                    st.success(f"✅ تم استخراج {len(contract_text)} حرف")
            else:
                st.error("❌ مكتبة python-docx غير مثبتة")

    elif input_method == "🖼️ صورة":
        img_file = st.file_uploader("ارفع صورة العقد", type=["jpg", "jpeg", "png", "webp"])
        if img_file:
            file_bytes = img_file.read()
            ext = img_file.name.split(".")[-1].lower()
            mime = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png", "webp": "image/webp"}.get(ext, "image/jpeg")
            image_data_dict = {"media_type": mime, "data": file_to_base64(file_bytes)}
            st.image(img_file, caption="صورة العقد", use_container_width=True)

    can_analyze = bool(contract_text.strip()) or image_data_dict is not None

    if can_analyze:
        st.markdown("<div class='section-divider'></div>", unsafe_allow_html=True)
        if st.button("🔍 حلل العقد الآن"):
            if not api_key:
                st.error("❌ أدخل Gemini API Key أولاً من الشريط الجانبي")
            else:
                with st.spinner("🤖 الذكاء الاصطناعي يحلل العقد..."):
                    try:
                        opts = {
                            "check_missing": check_missing,
                            "check_risks": check_risks,
                            "suggest_improvements": suggest_improvements,
                            "generate_score": generate_score,
                            "detail_level": detail_level,
                        }
                        result = analyze_contract_gemini(
                            api_key, contract_text,
                            contract_type if contract_type != "تلقائي (الذكاء يحدده)" else "غير محدد",
                            opts, image_data_dict
                        )
                        st.session_state["analysis_result"] = result
                    except json.JSONDecodeError as e:
                        st.error(f"❌ خطأ في قراءة الرد — جرب مرة أخرى: {e}")
                    except Exception as e:
                        st.error(f"❌ خطأ: {e}")

    if "analysis_result" in st.session_state:
        r = st.session_state["analysis_result"]
        st.markdown("---")
        st.markdown("## 📊 نتائج التحليل")

        col_summary, col_score = st.columns([2, 1])
        with col_summary:
            st.markdown(f"""
<div class="analysis-card">
    <div class="card-title">📋 {r.get('contract_type','عقد')}</div>
    <p style="color:#c0c0d8;line-height:1.8">{r.get('summary','')}</p>
    <div style="margin-top:1rem">
        {''.join(f'<span class="tag tag-yellow">{p}</span>' for p in r.get('parties',[]))}
    </div>
</div>""", unsafe_allow_html=True)

        with col_score:
            if "score" in r:
                sc = r["score"]
                overall = sc.get("overall", 0)
                cls = score_class(overall)
                st.markdown(f"""
<div class="analysis-card" style="text-align:center">
    <div class="score-ring {cls}">{overall}</div>
    <p style="color:#a0a0c0;font-size:0.9rem">من 100</p>
    <small style="color:#a0a0c0">وضوح: {sc.get('clarity',0)} | اكتمال: {sc.get('completeness',0)}</small><br>
    <small style="color:#a0a0c0">حماية: {sc.get('legal_protection',0)} | توازن: {sc.get('balance',0)}</small>
</div>""", unsafe_allow_html=True)

        verdict = r.get("verdict", "")
        v_color = "#00d4aa" if "جيد" in verdict else ("#f7c948" if "تعديل" in verdict else "#e94560")
        st.markdown(f"""
<div class="analysis-card" style="border-right:4px solid {v_color}">
    <div class="card-title" style="color:{v_color}">⚖️ الحكم النهائي</div>
    <p style="color:#e8e8f0;font-size:1.1rem;font-weight:600">{verdict}</p>
    <p style="color:#a0a0c0;margin-top:0.5rem">{r.get('recommendation','')}</p>
</div>""", unsafe_allow_html=True)

        dt1, dt2, dt3, dt4, dt5 = st.tabs(["✅ نقاط القوة", "❌ نقاط الضعف", "🔍 البنود الناقصة", "⚠️ المخاطر", "💡 التحسينات"])

        with dt1:
            items = r.get("strengths", [])
            if not items: st.info("لم يتم تحديد نقاط قوة")
            for item in items:
                st.markdown(f"""
<div class="analysis-card strength-card">
    <div class="card-title strength-title">✅ {item.get('title','')}</div>
    <p style="color:#c0c0d8">{item.get('detail','')}</p>
</div>""", unsafe_allow_html=True)

        with dt2:
            items = r.get("weaknesses", [])
            if not items: st.info("لم يتم تحديد نقاط ضعف")
            for item in items:
                st.markdown(f"""
<div class="analysis-card weakness-card">
    <div class="card-title weakness-title">❌ {item.get('title','')}</div>
    <p style="color:#c0c0d8">{item.get('detail','')}</p>
</div>""", unsafe_allow_html=True)

        with dt3:
            items = r.get("missing_clauses", [])
            if not items: st.info("لا توجد بنود ناقصة")
            for item in items:
                imp = item.get("importance", "متوسط")
                imp_color = "#e94560" if imp == "عالي" else ("#f7c948" if imp == "متوسط" else "#a0a0c0")
                st.markdown(f"""
<div class="analysis-card missing-card">
    <div class="card-title missing-title">🔍 {item.get('title','')} <span style="font-size:0.8rem;color:{imp_color}">({imp})</span></div>
    <p style="color:#c0c0d8"><strong style="color:#ff6b35">نص مقترح:</strong> {item.get('suggestion','')}</p>
</div>""", unsafe_allow_html=True)

        with dt4:
            items = r.get("risks", [])
            if not items: st.info("لا توجد مخاطر محددة")
            for item in items:
                sev = item.get("severity", "متوسط")
                sev_color = "#e94560" if sev == "عالي" else ("#f7c948" if sev == "متوسط" else "#a0a0c0")
                st.markdown(f"""
<div class="analysis-card" style="border-right:4px solid {sev_color}">
    <div class="card-title" style="color:{sev_color}">⚠️ {item.get('title','')} <span style="font-size:0.8rem">— خطورة: {sev}</span></div>
    <p style="color:#c0c0d8"><strong style="color:#f7c948">التخفيف:</strong> {item.get('mitigation','')}</p>
</div>""", unsafe_allow_html=True)

        with dt5:
            items = r.get("improvements", [])
            if not items: st.info("لا توجد تحسينات مقترحة")
            for item in items:
                pri = item.get("priority", "متوسط")
                pri_color = "#e94560" if pri == "عالي" else ("#f7c948" if pri == "متوسط" else "#00d4aa")
                st.markdown(f"""
<div class="analysis-card suggestion-card">
    <div class="card-title suggestion-title">💡 {item.get('title','')} <span style="font-size:0.8rem;color:{pri_color}">— أولوية: {pri}</span></div>
    <p style="color:#c0c0d8">{item.get('detail','')}</p>
</div>""", unsafe_allow_html=True)

        st.markdown("---")
        st.download_button(
            "⬇️ تحميل التقرير (JSON)",
            data=json.dumps(r, ensure_ascii=False, indent=2),
            file_name="contract_analysis.json",
            mime="application/json",
        )


# ══════════════════════════════════════════════
# TAB 2 – GENERATE
# ══════════════════════════════════════════════
with tab2:
    st.markdown("### ✍️ إنشاء عقد جديد")
    gen_type = st.selectbox("نوع العقد المطلوب", [
        "عقد عمل", "عقد إيجار سكني", "عقد إيجار تجاري",
        "عقد بيع عقار", "عقد بيع سيارة", "عقد شراكة تجارية",
        "عقد خدمات استشارية", "عقد توريد بضائع",
        "عقد مقاولة", "عقد نقل ملكية", "عقد قرض", "آخر",
    ])
    gen_details = st.text_area("تفاصيل العقد", height=200,
        placeholder="مثال:\nالطرف الأول: شركة ABC\nالطرف الثاني: محمد أحمد\nالراتب: 8000 جنيه\nالمدة: سنة...")

    if st.button("🪄 أنشئ العقد الآن"):
        if not api_key:
            st.error("❌ أدخل API Key أولاً")
        elif not gen_details.strip():
            st.warning("⚠️ أدخل تفاصيل العقد")
        else:
            with st.spinner("✍️ جاري كتابة العقد..."):
                try:
                    generated = generate_contract_gemini(api_key, gen_type, gen_details)
                    st.session_state["generated_contract"] = generated
                except Exception as e:
                    st.error(f"❌ خطأ: {e}")

    if "generated_contract" in st.session_state:
        st.markdown("---")
        st.markdown("### 📄 العقد المُنشأ")
        st.markdown(f'<div class="contract-output">{st.session_state["generated_contract"]}</div>',
                    unsafe_allow_html=True)
        st.download_button(
            "⬇️ تحميل العقد (TXT)",
            data=st.session_state["generated_contract"],
            file_name=f"{gen_type}.txt",
            mime="text/plain",
        )


# ══════════════════════════════════════════════
# TAB 3 – GUIDE
# ══════════════════════════════════════════════
with tab3:
    st.markdown("### 📚 دليل البنود الأساسية للعقود")
    clauses = {
        "البنود الأساسية في كل عقد": {
            "color": "#00d4aa",
            "items": [
                ("تعريف الأطراف", "الاسم الكامل، العنوان، رقم الهوية/السجل التجاري"),
                ("موضوع العقد", "وصف واضح ودقيق لما يتم الاتفاق عليه"),
                ("المقابل المادي", "المبلغ، طريقة الدفع، المواعيد"),
                ("مدة العقد", "تاريخ البداية والنهاية أو شروط الاستمرار"),
                ("التزامات الأطراف", "ما يجب على كل طرف فعله أو الامتناع عنه"),
            ]
        },
        "بنود الحماية القانونية": {
            "color": "#f7c948",
            "items": [
                ("السرية وعدم الإفصاح", "حماية المعلومات الحساسة"),
                ("الملكية الفكرية", "تحديد من يملك الإنتاج الفكري"),
                ("التعويضات والضمانات", "ما يحدث عند الإخلال بالعقد"),
                ("القوة القاهرة", "الأحداث الاستثنائية خارج السيطرة"),
                ("التنازل عن العقد", "شروط نقل الحقوق لطرف ثالث"),
            ]
        },
        "بنود إنهاء العقد": {
            "color": "#e94560",
            "items": [
                ("شروط الفسخ", "الأسباب التي تسمح بإنهاء العقد"),
                ("الإشعار المسبق", "المدة المطلوبة قبل الإنهاء"),
                ("التبعات عند الإنهاء", "ماذا يحدث للالتزامات بعد الإنهاء"),
                ("استرداد الأصول", "إعادة الممتلكات أو المعلومات"),
                ("التزامات ما بعد الإنهاء", "البنود التي تستمر بعد انتهاء العقد"),
            ]
        },
        "بنود حل النزاعات": {
            "color": "#a78bfa",
            "items": [
                ("القانون الحاكم", "أي قانون يطبق على العقد"),
                ("الاختصاص القضائي", "أي محكمة تختص بالنزاعات"),
                ("التحكيم", "هل يتم اللجوء للتحكيم بدل المحاكم"),
                ("الوساطة", "مرحلة الوساطة قبل التقاضي"),
                ("اللغة الرسمية", "اللغة المعتمدة في تفسير العقد"),
            ]
        },
    }
    for section, data in clauses.items():
        with st.expander(f"📌 {section}"):
            for title, desc in data["items"]:
                st.markdown(f"""
<div class="analysis-card" style="border-right:3px solid {data['color']};padding:1rem">
    <strong style="color:{data['color']}">{title}</strong>
    <p style="color:#a0a0c0;margin:0.3rem 0 0;font-size:0.95rem">{desc}</p>
</div>""", unsafe_allow_html=True)

# Footer
st.markdown("""
<div style="text-align:center;padding:2rem;color:#404060;font-size:0.85rem;
            margin-top:3rem;border-top:1px solid #2a2a4a">
    ⚖️ محلل العقود الذكي | مدعوم بـ Google Gemini AI — مجاني 100%<br>
    <small>هذا التطبيق للأغراض المعلوماتية فقط ولا يُغني عن استشارة محامٍ متخصص</small>
</div>
""", unsafe_allow_html=True)
try:
    import fitz
except ImportError:
    install("pymupdf")
    import fitz

try:
    from docx import Document
except ImportError:
    install("python-docx")
    from docx import Document

import streamlit as st
import base64
import json
import re
import io

PYMUPDF_AVAILABLE = True
DOCX_AVAILABLE = True

# ─────────────────────────────────────────────
# Page config
# ─────────────────────────────────────────────
st.set_page_config(
    page_title="محلل العقود الذكي",
    page_icon="⚖️",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ─────────────────────────────────────────────
# Custom CSS
# ─────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Tajawal:wght@300;400;500;700;800&display=swap');

* { font-family: 'Tajawal', sans-serif !important; }

html, body, [class*="css"] {
    direction: rtl;
    background: #0f0f1a;
    color: #e8e8f0;
}

.stApp { background: #0f0f1a; }

.main-header {
    background: linear-gradient(135deg, #1a1a2e 0%, #16213e 50%, #0f3460 100%);
    border: 1px solid #e94560;
    border-radius: 16px;
    padding: 2rem;
    text-align: center;
    margin-bottom: 2rem;
}
.main-header h1 { color: #e94560; font-size: 2.5rem; font-weight: 800; margin: 0; text-shadow: 0 0 30px rgba(233,69,96,0.5); }
.main-header p  { color: #a0a0c0; font-size: 1.1rem; margin: 0.5rem 0 0; }

.analysis-card {
    background: linear-gradient(135deg, #1a1a2e, #16213e);
    border: 1px solid #2a2a4a;
    border-radius: 12px;
    padding: 1.5rem;
    margin-bottom: 1.2rem;
}
.strength-card  { border-right: 4px solid #00d4aa; background: linear-gradient(135deg,#0d2b25,#0f2d27); }
.weakness-card  { border-right: 4px solid #e94560; background: linear-gradient(135deg,#2b0d1a,#2d0f1e); }
.suggestion-card{ border-right: 4px solid #f7c948; background: linear-gradient(135deg,#2b2500,#2d2600); }
.missing-card   { border-right: 4px solid #ff6b35; background: linear-gradient(135deg,#2b1500,#2d1600); }

.card-title { font-size: 1.2rem; font-weight: 700; margin-bottom: 1rem; display: flex; align-items: center; gap: 0.5rem; }
.strength-title  { color: #00d4aa; }
.weakness-title  { color: #e94560; }
.suggestion-title{ color: #f7c948; }
.missing-title   { color: #ff6b35; }

.score-ring {
    width: 120px; height: 120px; border-radius: 50%;
    display: flex; align-items: center; justify-content: center;
    margin: 0 auto 1rem; font-size: 2rem; font-weight: 800; border: 6px solid;
}
.score-high { border-color:#00d4aa; color:#00d4aa; background:rgba(0,212,170,0.1); }
.score-mid  { border-color:#f7c948; color:#f7c948; background:rgba(247,201,72,0.1); }
.score-low  { border-color:#e94560; color:#e94560; background:rgba(233,69,96,0.1); }

.stTabs [data-baseweb="tab-list"] { background:#1a1a2e; border-radius:10px; padding:4px; gap:4px; }
.stTabs [data-baseweb="tab"] { background:transparent; border-radius:8px; color:#a0a0c0; font-weight:600; padding:8px 20px; }
.stTabs [aria-selected="true"] { background:#e94560 !important; color:white !important; }

.stButton > button {
    background: linear-gradient(135deg,#e94560,#c73652);
    color: white; border: none; border-radius: 10px;
    padding: 0.7rem 2rem; font-size: 1.1rem; font-weight: 700; width: 100%;
}
.stButton > button:hover { transform: translateY(-2px); box-shadow: 0 8px 25px rgba(233,69,96,0.4); }

.stTextArea textarea, .stTextInput input {
    background:#1a1a2e !important; border:1px solid #2a2a4a !important;
    border-radius:10px !important; color:#e8e8f0 !important; direction:rtl;
}
.stTextArea textarea:focus, .stTextInput input:focus {
    border-color:#e94560 !important; box-shadow:0 0 0 2px rgba(233,69,96,0.2) !important;
}

[data-testid="stSidebar"] { background:#12121f !important; border-left:1px solid #2a2a4a; }

.tag { display:inline-block; padding:3px 12px; border-radius:20px; font-size:0.85rem; font-weight:600; margin:3px; }
.tag-green  { background:rgba(0,212,170,0.2); color:#00d4aa; border:1px solid #00d4aa; }
.tag-red    { background:rgba(233,69,96,0.2);  color:#e94560; border:1px solid #e94560; }
.tag-yellow { background:rgba(247,201,72,0.2); color:#f7c948; border:1px solid #f7c948; }

.section-divider { height:2px; background:linear-gradient(90deg,#e94560,transparent); margin:1.5rem 0; border-radius:2px; }

.contract-output {
    background:#1a1a2e; border:1px solid #2a2a4a; border-radius:12px;
    padding:2rem; white-space:pre-wrap; line-height:2; direction:rtl;
}
</style>
""", unsafe_allow_html=True)

# ─────────────────────────────────────────────
# Header
# ─────────────────────────────────────────────
st.markdown("""
<div class="main-header">
    <h1>⚖️ محلل العقود الذكي</h1>
    <p>تحليل شامل واحترافي للعقود القانونية بالذكاء الاصطناعي — مجاني 100%</p>
</div>
""", unsafe_allow_html=True)

# ─────────────────────────────────────────────
# Sidebar
# ─────────────────────────────────────────────
with st.sidebar:
    st.markdown("### 🔑 إعدادات API")
    st.markdown("""
<div style="background:rgba(0,212,170,0.1);border:1px solid #00d4aa;border-radius:8px;padding:0.8rem;margin-bottom:1rem;font-size:0.85rem;color:#00d4aa">
    ✅ Google Gemini مجاني تماماً!<br>
    احصل على مفتاحك من:<br>
    <a href="https://aistudio.google.com" target="_blank" style="color:#00d4aa">aistudio.google.com</a>
</div>
""", unsafe_allow_html=True)
    api_key = st.text_input("Gemini API Key", type="password", placeholder="AIza...")

    st.markdown("---")
    st.markdown("### 📋 نوع العقد")
    contract_type = st.selectbox("اختر نوع العقد", [
        "تلقائي (الذكاء يحدده)", "عقد عمل", "عقد إيجار", "عقد بيع",
        "عقد شراكة", "عقد خدمات", "عقد توريد", "عقد استشارات",
        "عقد نقل ملكية", "عقد قرض", "أخرى",
    ])

    st.markdown("---")
    st.markdown("### ⚙️ خيارات التحليل")
    check_missing        = st.checkbox("🔍 تحقق من البنود الناقصة", value=True)
    check_risks          = st.checkbox("⚠️ تحليل المخاطر القانونية", value=True)
    suggest_improvements = st.checkbox("💡 اقتراح تحسينات", value=True)
    generate_score       = st.checkbox("🎯 تقييم شامل (نقاط)", value=True)

    st.markdown("---")
    st.markdown("### 📊 مستوى التفصيل")
    detail_level = st.radio("", ["مختصر", "متوسط", "تفصيلي مكثف"], index=1)


# ─────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────
def extract_text_from_pdf(file_bytes: bytes) -> str:
    if not PYMUPDF_AVAILABLE:
        return ""
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    return "\n".join(page.get_text("text") for page in doc).strip()

def extract_text_from_docx(file_bytes: bytes) -> str:
    if not DOCX_AVAILABLE:
        return ""
    doc = Document(io.BytesIO(file_bytes))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

def file_to_base64(file_bytes: bytes) -> str:
    return base64.standard_b64encode(file_bytes).decode("utf-8")

def score_class(s):
    return "score-high" if s >= 70 else ("score-mid" if s >= 50 else "score-low")


# ─────────────────────────────────────────────
# Gemini: Analyze contract
# ─────────────────────────────────────────────
def analyze_contract_gemini(api_key: str, contract_text: str,
                             contract_type_hint: str, options: dict,
                             image_data: dict | None = None) -> dict:

    genai.configure(api_key=api_key)
    model = genai.GenerativeModel("gemini-2.0-flash")

    detail_map = {"مختصر": "موجز", "متوسط": "متوسط التفصيل", "تفصيلي مكثف": "تفصيلي جداً"}
    detail = detail_map.get(options["detail_level"], "متوسط التفصيل")

    options_text = []
    if options["check_missing"]:        options_text.append("- تحقق من البنود الناقصة")
    if options["check_risks"]:          options_text.append("- حلل المخاطر القانونية")
    if options["suggest_improvements"]: options_text.append("- قدم اقتراحات تحسين")
    if options["generate_score"]:       options_text.append("- قيّم العقد بنقاط من 100")

    prompt = f"""أنت محامٍ وخبير قانوني متخصص في تحليل العقود، تتحدث العربية بطلاقة.
أجب فقط بـ JSON صحيح بدون أي نص إضافي أو ```json أو أي شيء آخر.

حلل العقد التالي وأعطني JSON بهذا التنسيق بالضبط:

{{
  "contract_type": "نوع العقد",
  "summary": "ملخص العقد في 3-5 جمل",
  "parties": ["الطرف الأول: ...", "الطرف الثاني: ..."],
  "key_terms": ["بند مهم 1", "بند مهم 2"],
  "strengths": [
    {{"title": "عنوان نقطة القوة", "detail": "شرح تفصيلي"}}
  ],
  "weaknesses": [
    {{"title": "عنوان نقطة الضعف", "detail": "شرح تفصيلي"}}
  ],
  "missing_clauses": [
    {{"title": "البند الناقص", "importance": "عالي", "suggestion": "نص مقترح"}}
  ],
  "risks": [
    {{"title": "الخطر", "severity": "عالي", "mitigation": "طريقة التخفيف"}}
  ],
  "improvements": [
    {{"title": "التحسين", "detail": "شرح التحسين", "priority": "عالي"}}
  ],
  "score": {{
    "overall": 75,
    "clarity": 80,
    "completeness": 70,
    "legal_protection": 75,
    "balance": 65,
    "justification": "سبب التقييم"
  }},
  "verdict": "الحكم النهائي على العقد",
  "recommendation": "توصية مفصلة للشخص قبل التوقيع"
}}

نوع العقد: {contract_type_hint}
مستوى التفصيل: {detail}
خيارات التحليل:
{chr(10).join(options_text)}

نص العقد:
{contract_text[:12000]}
"""

    if image_data:
        image_part = {
            "mime_type": image_data["media_type"],
            "data": base64.b64decode(image_data["data"])
        }
        response = model.generate_content([prompt, image_part])
    else:
        response = model.generate_content(prompt)

    raw = response.text.strip()
    raw = re.sub(r"^```(?:json)?\s*", "", raw)
    raw = re.sub(r"\s*```$", "", raw)
    return json.loads(raw)


# ─────────────────────────────────────────────
# Gemini: Generate contract
# ─────────────────────────────────────────────
def generate_contract_gemini(api_key: str, contract_type: str, details: str) -> str:
    genai.configure(api_key=api_key)
    model = genai.GenerativeModel("gemini-2.0-flash")

    prompt = f"""أنت محامٍ خبير في صياغة العقود القانونية باللغة العربية.
اكتب عقد {contract_type} احترافي وشامل باللغة العربية.

التفاصيل:
{details}

اكتب العقد بالكامل مع:
- ديباجة العقد وأطرافه
- جميع البنود الأساسية والفرعية
- الشروط والأحكام
- بنود الضمانات والتعويضات
- بنود فسخ العقد
- القانون الحاكم وحل النزاعات
- التواقيع والتوثيق

اجعل العقد متوازناً ويحمي حقوق الطرفين."""

    response = model.generate_content(prompt)
    return response.text


# ─────────────────────────────────────────────
# Main Tabs
# ─────────────────────────────────────────────
tab1, tab2, tab3 = st.tabs(["📄 تحليل عقد موجود", "✍️ إنشاء عقد جديد", "📚 دليل البنود"])

# ══════════════════════════════════════════════
# TAB 1 – ANALYZE
# ══════════════════════════════════════════════
with tab1:
    st.markdown("### 📤 رفع أو كتابة العقد")
    input_method = st.radio("طريقة الإدخال", ["📝 نص مباشر", "📎 PDF", "📄 Word", "🖼️ صورة"], horizontal=True)

    contract_text = ""
    image_data_dict = None

    if input_method == "📝 نص مباشر":
        contract_text = st.text_area("الصق نص العقد هنا", height=300,
                                     placeholder="اكتب أو الصق نص العقد هنا...")

    elif input_method == "📎 PDF":
        pdf_file = st.file_uploader("ارفع ملف PDF", type=["pdf"])
        if pdf_file:
            file_bytes = pdf_file.read()
            if PYMUPDF_AVAILABLE:
                with st.spinner("⏳ جاري استخراج النص..."):
                    contract_text = extract_text_from_pdf(file_bytes)
                if contract_text:
                    st.success(f"✅ تم استخراج {len(contract_text)} حرف")
                    with st.expander("👁️ عرض النص المستخرج"):
                        st.text(contract_text[:3000] + ("..." if len(contract_text) > 3000 else ""))
                else:
                    st.warning("⚠️ لم يُستخرج نص — سيُرسل كصورة")
                    image_data_dict = {"media_type": "application/pdf", "data": file_to_base64(file_bytes)}
            else:
                image_data_dict = {"media_type": "application/pdf", "data": file_to_base64(file_bytes)}
                st.info("📤 سيُرسل PDF للذكاء الاصطناعي مباشرة")

    elif input_method == "📄 Word":
        docx_file = st.file_uploader("ارفع ملف Word", type=["docx"])
        if docx_file:
            file_bytes = docx_file.read()
            if DOCX_AVAILABLE:
                with st.spinner("⏳ جاري استخراج النص..."):
                    contract_text = extract_text_from_docx(file_bytes)
                if contract_text:
                    st.success(f"✅ تم استخراج {len(contract_text)} حرف")
            else:
                st.error("❌ مكتبة python-docx غير مثبتة")

    elif input_method == "🖼️ صورة":
        img_file = st.file_uploader("ارفع صورة العقد", type=["jpg", "jpeg", "png", "webp"])
        if img_file:
            file_bytes = img_file.read()
            ext = img_file.name.split(".")[-1].lower()
            mime = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png", "webp": "image/webp"}.get(ext, "image/jpeg")
            image_data_dict = {"media_type": mime, "data": file_to_base64(file_bytes)}
            st.image(img_file, caption="صورة العقد", use_container_width=True)

    can_analyze = bool(contract_text.strip()) or image_data_dict is not None

    if can_analyze:
        st.markdown("<div class='section-divider'></div>", unsafe_allow_html=True)
        if st.button("🔍 حلل العقد الآن"):
            if not api_key:
                st.error("❌ أدخل Gemini API Key أولاً من الشريط الجانبي")
            else:
                with st.spinner("🤖 الذكاء الاصطناعي يحلل العقد..."):
                    try:
                        opts = {
                            "check_missing": check_missing,
                            "check_risks": check_risks,
                            "suggest_improvements": suggest_improvements,
                            "generate_score": generate_score,
                            "detail_level": detail_level,
                        }
                        result = analyze_contract_gemini(
                            api_key, contract_text,
                            contract_type if contract_type != "تلقائي (الذكاء يحدده)" else "غير محدد",
                            opts, image_data_dict
                        )
                        st.session_state["analysis_result"] = result
                    except json.JSONDecodeError as e:
                        st.error(f"❌ خطأ في قراءة الرد — جرب مرة أخرى: {e}")
                    except Exception as e:
                        st.error(f"❌ خطأ: {e}")

    if "analysis_result" in st.session_state:
        r = st.session_state["analysis_result"]
        st.markdown("---")
        st.markdown("## 📊 نتائج التحليل")

        col_summary, col_score = st.columns([2, 1])
        with col_summary:
            st.markdown(f"""
<div class="analysis-card">
    <div class="card-title">📋 {r.get('contract_type','عقد')}</div>
    <p style="color:#c0c0d8;line-height:1.8">{r.get('summary','')}</p>
    <div style="margin-top:1rem">
        {''.join(f'<span class="tag tag-yellow">{p}</span>' for p in r.get('parties',[]))}
    </div>
</div>""", unsafe_allow_html=True)

        with col_score:
            if "score" in r:
                sc = r["score"]
                overall = sc.get("overall", 0)
                cls = score_class(overall)
                st.markdown(f"""
<div class="analysis-card" style="text-align:center">
    <div class="score-ring {cls}">{overall}</div>
    <p style="color:#a0a0c0;font-size:0.9rem">من 100</p>
    <small style="color:#a0a0c0">وضوح: {sc.get('clarity',0)} | اكتمال: {sc.get('completeness',0)}</small><br>
    <small style="color:#a0a0c0">حماية: {sc.get('legal_protection',0)} | توازن: {sc.get('balance',0)}</small>
</div>""", unsafe_allow_html=True)

        verdict = r.get("verdict", "")
        v_color = "#00d4aa" if "جيد" in verdict else ("#f7c948" if "تعديل" in verdict else "#e94560")
        st.markdown(f"""
<div class="analysis-card" style="border-right:4px solid {v_color}">
    <div class="card-title" style="color:{v_color}">⚖️ الحكم النهائي</div>
    <p style="color:#e8e8f0;font-size:1.1rem;font-weight:600">{verdict}</p>
    <p style="color:#a0a0c0;margin-top:0.5rem">{r.get('recommendation','')}</p>
</div>""", unsafe_allow_html=True)

        dt1, dt2, dt3, dt4, dt5 = st.tabs(["✅ نقاط القوة", "❌ نقاط الضعف", "🔍 البنود الناقصة", "⚠️ المخاطر", "💡 التحسينات"])

        with dt1:
            items = r.get("strengths", [])
            if not items: st.info("لم يتم تحديد نقاط قوة")
            for item in items:
                st.markdown(f"""
<div class="analysis-card strength-card">
    <div class="card-title strength-title">✅ {item.get('title','')}</div>
    <p style="color:#c0c0d8">{item.get('detail','')}</p>
</div>""", unsafe_allow_html=True)

        with dt2:
            items = r.get("weaknesses", [])
            if not items: st.info("لم يتم تحديد نقاط ضعف")
            for item in items:
                st.markdown(f"""
<div class="analysis-card weakness-card">
    <div class="card-title weakness-title">❌ {item.get('title','')}</div>
    <p style="color:#c0c0d8">{item.get('detail','')}</p>
</div>""", unsafe_allow_html=True)

        with dt3:
            items = r.get("missing_clauses", [])
            if not items: st.info("لا توجد بنود ناقصة")
            for item in items:
                imp = item.get("importance", "متوسط")
                imp_color = "#e94560" if imp == "عالي" else ("#f7c948" if imp == "متوسط" else "#a0a0c0")
                st.markdown(f"""
<div class="analysis-card missing-card">
    <div class="card-title missing-title">🔍 {item.get('title','')} <span style="font-size:0.8rem;color:{imp_color}">({imp})</span></div>
    <p style="color:#c0c0d8"><strong style="color:#ff6b35">نص مقترح:</strong> {item.get('suggestion','')}</p>
</div>""", unsafe_allow_html=True)

        with dt4:
            items = r.get("risks", [])
            if not items: st.info("لا توجد مخاطر محددة")
            for item in items:
                sev = item.get("severity", "متوسط")
                sev_color = "#e94560" if sev == "عالي" else ("#f7c948" if sev == "متوسط" else "#a0a0c0")
                st.markdown(f"""
<div class="analysis-card" style="border-right:4px solid {sev_color}">
    <div class="card-title" style="color:{sev_color}">⚠️ {item.get('title','')} <span style="font-size:0.8rem">— خطورة: {sev}</span></div>
    <p style="color:#c0c0d8"><strong style="color:#f7c948">التخفيف:</strong> {item.get('mitigation','')}</p>
</div>""", unsafe_allow_html=True)

        with dt5:
            items = r.get("improvements", [])
            if not items: st.info("لا توجد تحسينات مقترحة")
            for item in items:
                pri = item.get("priority", "متوسط")
                pri_color = "#e94560" if pri == "عالي" else ("#f7c948" if pri == "متوسط" else "#00d4aa")
                st.markdown(f"""
<div class="analysis-card suggestion-card">
    <div class="card-title suggestion-title">💡 {item.get('title','')} <span style="font-size:0.8rem;color:{pri_color}">— أولوية: {pri}</span></div>
    <p style="color:#c0c0d8">{item.get('detail','')}</p>
</div>""", unsafe_allow_html=True)

        st.markdown("---")
        st.download_button(
            "⬇️ تحميل التقرير (JSON)",
            data=json.dumps(r, ensure_ascii=False, indent=2),
            file_name="contract_analysis.json",
            mime="application/json",
        )


# ══════════════════════════════════════════════
# TAB 2 – GENERATE
# ══════════════════════════════════════════════
with tab2:
    st.markdown("### ✍️ إنشاء عقد جديد")
    gen_type = st.selectbox("نوع العقد المطلوب", [
        "عقد عمل", "عقد إيجار سكني", "عقد إيجار تجاري",
        "عقد بيع عقار", "عقد بيع سيارة", "عقد شراكة تجارية",
        "عقد خدمات استشارية", "عقد توريد بضائع",
        "عقد مقاولة", "عقد نقل ملكية", "عقد قرض", "آخر",
    ])
    gen_details = st.text_area("تفاصيل العقد", height=200,
        placeholder="مثال:\nالطرف الأول: شركة ABC\nالطرف الثاني: محمد أحمد\nالراتب: 8000 جنيه\nالمدة: سنة...")

    if st.button("🪄 أنشئ العقد الآن"):
        if not api_key:
            st.error("❌ أدخل API Key أولاً")
        elif not gen_details.strip():
            st.warning("⚠️ أدخل تفاصيل العقد")
        else:
            with st.spinner("✍️ جاري كتابة العقد..."):
                try:
                    generated = generate_contract_gemini(api_key, gen_type, gen_details)
                    st.session_state["generated_contract"] = generated
                except Exception as e:
                    st.error(f"❌ خطأ: {e}")

    if "generated_contract" in st.session_state:
        st.markdown("---")
        st.markdown("### 📄 العقد المُنشأ")
        st.markdown(f'<div class="contract-output">{st.session_state["generated_contract"]}</div>',
                    unsafe_allow_html=True)
        st.download_button(
            "⬇️ تحميل العقد (TXT)",
            data=st.session_state["generated_contract"],
            file_name=f"{gen_type}.txt",
            mime="text/plain",
        )


# ══════════════════════════════════════════════
# TAB 3 – GUIDE
# ══════════════════════════════════════════════
with tab3:
    st.markdown("### 📚 دليل البنود الأساسية للعقود")
    clauses = {
        "البنود الأساسية في كل عقد": {
            "color": "#00d4aa",
            "items": [
                ("تعريف الأطراف", "الاسم الكامل، العنوان، رقم الهوية/السجل التجاري"),
                ("موضوع العقد", "وصف واضح ودقيق لما يتم الاتفاق عليه"),
                ("المقابل المادي", "المبلغ، طريقة الدفع، المواعيد"),
                ("مدة العقد", "تاريخ البداية والنهاية أو شروط الاستمرار"),
                ("التزامات الأطراف", "ما يجب على كل طرف فعله أو الامتناع عنه"),
            ]
        },
        "بنود الحماية القانونية": {
            "color": "#f7c948",
            "items": [
                ("السرية وعدم الإفصاح", "حماية المعلومات الحساسة"),
                ("الملكية الفكرية", "تحديد من يملك الإنتاج الفكري"),
                ("التعويضات والضمانات", "ما يحدث عند الإخلال بالعقد"),
                ("القوة القاهرة", "الأحداث الاستثنائية خارج السيطرة"),
                ("التنازل عن العقد", "شروط نقل الحقوق لطرف ثالث"),
            ]
        },
        "بنود إنهاء العقد": {
            "color": "#e94560",
            "items": [
                ("شروط الفسخ", "الأسباب التي تسمح بإنهاء العقد"),
                ("الإشعار المسبق", "المدة المطلوبة قبل الإنهاء"),
                ("التبعات عند الإنهاء", "ماذا يحدث للالتزامات بعد الإنهاء"),
                ("استرداد الأصول", "إعادة الممتلكات أو المعلومات"),
                ("التزامات ما بعد الإنهاء", "البنود التي تستمر بعد انتهاء العقد"),
            ]
        },
        "بنود حل النزاعات": {
            "color": "#a78bfa",
            "items": [
                ("القانون الحاكم", "أي قانون يطبق على العقد"),
                ("الاختصاص القضائي", "أي محكمة تختص بالنزاعات"),
                ("التحكيم", "هل يتم اللجوء للتحكيم بدل المحاكم"),
                ("الوساطة", "مرحلة الوساطة قبل التقاضي"),
                ("اللغة الرسمية", "اللغة المعتمدة في تفسير العقد"),
            ]
        },
    }
    for section, data in clauses.items():
        with st.expander(f"📌 {section}"):
            for title, desc in data["items"]:
                st.markdown(f"""
<div class="analysis-card" style="border-right:3px solid {data['color']};padding:1rem">
    <strong style="color:{data['color']}">{title}</strong>
    <p style="color:#a0a0c0;margin:0.3rem 0 0;font-size:0.95rem">{desc}</p>
</div>""", unsafe_allow_html=True)

# Footer
st.markdown("""
<div style="text-align:center;padding:2rem;color:#404060;font-size:0.85rem;
            margin-top:3rem;border-top:1px solid #2a2a4a">
    ⚖️ محلل العقود الذكي | مدعوم بـ Google Gemini AI — مجاني 100%<br>
    <small>هذا التطبيق للأغراض المعلوماتية فقط ولا يُغني عن استشارة محامٍ متخصص</small>
</div>
""", unsafe_allow_html=True)
